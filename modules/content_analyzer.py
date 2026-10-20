
import re
import csv
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt", ".csv", ".pdf", ".docx", ".xlsx"
}

SENSITIVE_KEYWORDS = [
    "confidential",
    "salary",
    "employee details",
    "bank account",
    "account number",
    "password",
    "secret key",
    "personal information",
    "credit card",
    "aadhaar",
    "pan number",
    "medical record",
    "private key",
    "customer details",
    "income details",
]

PATTERNS = {
    "Email address": re.compile(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
    ),
    "Phone number": re.compile(
        r"(?<!\d)(?:\+91[\s-]?)?[6-9]\d{9}(?!\d)"
    ),
    "Possible Aadhaar number": re.compile(
        r"(?<!\d)\d{4}[\s-]\d{4}[\s-]\d{4}(?!\d)"
    ),
    "Possible payment card number": re.compile(
        r"(?<!\d)(?:\d[ -]?){13,19}(?!\d)"
    ),
}


def extract_content(file_path):
    """Extract readable text from supported document types."""

    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        return None, f"Unsupported file type: {extension or 'unknown'}"

    try:
        if extension == ".txt":
            text = path.read_text(
                encoding="utf-8", errors="replace"
            )

        elif extension == ".csv":
            rows = []
            with path.open(
                "r", encoding="utf-8-sig", errors="replace",
                newline=""
            ) as file:
                for row in csv.reader(file):
                    rows.append(" ".join(row))
            text = "\n".join(rows)

        elif extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))

            if reader.is_encrypted:
                return None, "Password-protected PDF cannot be analyzed."

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif extension == ".docx":
            from docx import Document
            document = Document(str(path))

            parts = [p.text for p in document.paragraphs]

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(parts)

        elif extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                str(path), read_only=True, data_only=True
            )
            parts = []

            try:
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        values = [
                            str(value) for value in row
                            if value is not None
                        ]
                        if values:
                            parts.append(" ".join(values))
            finally:
                workbook.close()

            text = "\n".join(parts)

        return text, None

    except Exception as error:
        return None, f"Could not read file: {error}"


def analyze_file_content(file_path):
    """Analyze actual file content for sensitive information."""

    text, error = extract_content(file_path)

    if error:
        return {
            "file": Path(file_path).name,
            "status": "Not Analyzed",
            "sensitive": None,
            "keyword_matches": [],
            "pattern_matches": [],
            "message": error,
        }

    if not text or not text.strip():
        return {
            "file": Path(file_path).name,
            "status": "Not Analyzed",
            "sensitive": None,
            "keyword_matches": [],
            "pattern_matches": [],
            "message": "No readable text found. The file may be empty, scanned, or image-based.",
        }

    content = text.lower()

    keyword_matches = [
        keyword for keyword in SENSITIVE_KEYWORDS
        if keyword in content
    ]

    pattern_matches = [
        name for name, pattern in PATTERNS.items()
        if pattern.search(text)
    ]

    sensitive = bool(keyword_matches or pattern_matches)

    reasons = []
    if keyword_matches:
        reasons.append("Sensitive terms found")
    if pattern_matches:
        reasons.append("Possible personal or financial data found")

    return {
        "file": Path(file_path).name,
        "status": "Analyzed",
        "sensitive": sensitive,
        "keyword_matches": keyword_matches,
        "pattern_matches": pattern_matches,
        "message": (
            "; ".join(reasons)
            if reasons else "No configured sensitive indicators found"
        ),
    }

